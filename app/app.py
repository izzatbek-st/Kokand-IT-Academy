from datetime import datetime
from pathlib import Path
from threading import Lock
import json
import os
import uuid

from flask import Flask, jsonify, request, send_from_directory
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
STATIC_DIR = PROJECT_DIR / 'static'

app = Flask(__name__, static_folder=str(STATIC_DIR), static_url_path='/static')
EXCEL_FILE = PROJECT_DIR / 'Lidlar.xlsx'
PENDING_LEADS_FILE = PROJECT_DIR / 'pending_leads.json'
TELEGRAM_OUTBOX_DIR = PROJECT_DIR / 'telegram_outbox'
HEADERS = ['Ism', 'Telefon', 'Kurs', 'Izoh', 'Yozilgan Sana']
LEAD_LOCK = Lock()


def load_environment():
    """Load local settings without an additional dependency."""
    env_file = BASE_DIR / '.env'
    if not env_file.exists():
        return

    for line in env_file.read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if not line or line.startswith('#') or '=' not in line:
            continue
        key, value = line.split('=', 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if not os.getenv(key):
            os.environ[key] = value


def ensure_excel_file():
    if EXCEL_FILE.exists():
        return

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Lidlar'
    for index, header in enumerate(HEADERS, start=1):
        worksheet[f'{get_column_letter(index)}1'] = header
    workbook.save(EXCEL_FILE)


def load_pending_leads():
    if not PENDING_LEADS_FILE.exists():
        return []
    try:
        leads = json.loads(PENDING_LEADS_FILE.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        app.logger.exception('Pending lead file could not be read')
        return []
    return leads if isinstance(leads, list) else []


def save_pending_leads(leads):
    temporary_file = PENDING_LEADS_FILE.with_suffix('.tmp')
    temporary_file.write_text(json.dumps(leads, ensure_ascii=False), encoding='utf-8')
    temporary_file.replace(PENDING_LEADS_FILE)


def save_leads_to_excel(leads):
    ensure_excel_file()
    workbook = load_workbook(EXCEL_FILE)
    worksheet = workbook.active
    for lead in leads:
        worksheet.append([
            lead['name'], lead['phone'], lead['course'], lead['note'], lead['submitted_at'],
        ])
    workbook.save(EXCEL_FILE)


def persist_lead(lead):
    """Save to Excel, or keep it in a local queue while Office has the file open."""
    with LEAD_LOCK:
        pending_leads = load_pending_leads()
        try:
            save_leads_to_excel([*pending_leads, lead])
        except PermissionError:
            save_pending_leads([*pending_leads, lead])
            return False

        if PENDING_LEADS_FILE.exists():
            PENDING_LEADS_FILE.unlink()
        return True


def queue_telegram_lead(lead):
    """The bot sends outbox records, so the web server does not need Telegram access."""
    TELEGRAM_OUTBOX_DIR.mkdir(exist_ok=True)
    message_id = uuid.uuid4().hex
    temporary_file = TELEGRAM_OUTBOX_DIR / f'{message_id}.tmp'
    message_file = TELEGRAM_OUTBOX_DIR / f'{message_id}.json'
    temporary_file.write_text(json.dumps(lead, ensure_ascii=False), encoding='utf-8')
    temporary_file.replace(message_file)


@app.route('/')
def index():
    return send_from_directory(STATIC_DIR, 'website1.html')


@app.route('/images/<filename>')
def serve_image(filename):
    return send_from_directory(STATIC_DIR / 'images', filename)


@app.route('/files/<filename>')
def serve_file_download(filename):
    return send_from_directory(STATIC_DIR / 'files', filename)


@app.route('/submit', methods=['POST'])
def submit_lead():
    data = request.get_json(silent=True)
    if not data:
        return jsonify({'success': False, 'message': 'Malumotlar topilmadi.'}), 400

    name = data.get('name', '').strip()
    phone = data.get('phone', '').strip()
    course = data.get('course', '').strip()
    note = data.get('note', '').strip()
    if not name or not phone or not course:
        return jsonify({
            'success': False,
            'message': 'Iltimos, ism, telefon va kursni kiriting.',
        }), 400

    lead = {
        'name': name,
        'phone': phone,
        'course': course,
        'note': note,
        'submitted_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
    }

    try:
        excel_saved = persist_lead(lead)
        queue_telegram_lead(lead)
    except (OSError, ValueError, KeyError):
        app.logger.exception('Lead could not be saved')
        return jsonify({
            'success': False,
            'message': 'Sorovni saqlashda xatolik yuz berdi. Qayta urinib koring.',
        }), 500

    if excel_saved:
        message = 'Sorov qabul qilindi. Admin botiga yuboriladi.'
    else:
        message = 'Sorov qabul qilindi. Excel fayli ochiq, shuning uchun royxatga keyin qoshiladi.'
    return jsonify({'success': True, 'message': message})


if __name__ == '__main__':
    app.run(port=5000)
